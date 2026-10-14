"""
Génère la présentation PowerPoint du projet (16:9, 14 diapositives, notes d'orateur) :
documents/rapports/Presentation_NURU_HR_Portfolio.pptx

Chaque forme est aussi dessinée dans un aperçu PNG (documents/rapports/apercu_diapositives/)
pour contrôler visuellement la mise en page, et chaque texte est mesuré pour détecter
tout débordement.

Utilisation (depuis la racine du projet) :
    python documents/rapports/generer_presentation.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import contenu_rapport as cr  # noqa: E402
from PIL import Image, ImageDraw, ImageFont  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

C = cr.COULEURS
FICHIER = cr.RAPPORTS / "Presentation_NURU_HR_Portfolio.pptx"
APERCU = cr.RAPPORTS / "apercu_diapositives"
L, H = 13.333, 7.5          # pouces (16:9)
ECHELLE = 96                # pixels par pouce dans l'aperçu
POLICES = Path("C:/Windows/Fonts")
FONT = {False: str(POLICES / "calibri.ttf"), True: str(POLICES / "calibrib.ttf")}


class Diapo:
    """Ajoute chaque forme au fichier PowerPoint ET à l'aperçu, et mesure les textes."""

    def __init__(self, deck, fond=C["blanc"], numero=None):
        self.deck = deck
        self.slide = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
        self.img = Image.new("RGB", (int(L * ECHELLE), int(H * ECHELLE)), "#" + fond)
        self.draw = ImageDraw.Draw(self.img)
        self.fond = fond
        bg = self.slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor.from_string(fond)
        if numero:
            self.texte(0.6, 7.0, 9, 0.3, "Données fictives · projet personnel indépendant de Nuru", 10, C["gris"])
            self.texte(L - 1.6, 7.0, 1.0, 0.3, str(numero), 10, C["gris"], align="r")

    # --- formes -------------------------------------------------------------------------------
    def forme(self, x, y, w, h, couleur, type_="rect", ligne=None, nom=None):
        typ = {"rect": MSO_SHAPE.RECTANGLE, "arrondi": MSO_SHAPE.ROUNDED_RECTANGLE, "cercle": MSO_SHAPE.OVAL}[type_]
        s = self.slide.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
        if nom:
            s.name = nom
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(couleur)
        if ligne:
            s.line.color.rgb = RGBColor.from_string(ligne)
            s.line.width = Pt(1)
        else:
            s.line.fill.background()
        s.shadow.inherit = False
        if type_ == "arrondi":
            s.adjustments[0] = 0.08
        box = [int(x * ECHELLE), int(y * ECHELLE), int((x + w) * ECHELLE), int((y + h) * ECHELLE)]
        trait = "#" + ligne if ligne else None
        if type_ == "cercle":
            self.draw.ellipse(box, fill="#" + couleur, outline=trait)
        elif type_ == "arrondi":
            self.draw.rounded_rectangle(box, radius=int(min(w, h) * 0.08 * ECHELLE), fill="#" + couleur, outline=trait)
        else:
            self.draw.rectangle(box, fill="#" + couleur, outline=trait)
        return s

    def trait(self, x1, y1, x2, y2, couleur, epaisseur=2):
        c = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = RGBColor.from_string(couleur)
        c.line.width = Pt(epaisseur)
        self.draw.line([x1 * ECHELLE, y1 * ECHELLE, x2 * ECHELLE, y2 * ECHELLE], fill="#" + couleur,
                       width=max(1, int(epaisseur * ECHELLE / 72)))

    def image(self, chemin, x, y, w, h, cadre=C["trait"]):
        """Image de largeur w, recadrée par le bas pour tenir dans la hauteur h (proportions conservées)."""
        with Image.open(chemin) as im:
            hauteur_utile = min(im.height, round(im.width * h / w))
            h = w * hauteur_utile / im.width          # image plus courte que la zone : on réduit h
            recadree = im.crop((0, 0, im.width, hauteur_utile)).convert("RGB")
            part_coupee = 1 - hauteur_utile / im.height
        pic = self.slide.shapes.add_picture(str(chemin), Inches(x), Inches(y), Inches(w), Inches(h))
        pic.crop_bottom = part_coupee
        pic.line.color.rgb = RGBColor.from_string(cadre)
        pic.line.width = Pt(1)
        self.img.paste(recadree.resize((int(w * ECHELLE), int(h * ECHELLE))), (int(x * ECHELLE), int(y * ECHELLE)))
        self.draw.rectangle([x * ECHELLE, y * ECHELLE, (x + w) * ECHELLE, (y + h) * ECHELLE], outline="#" + cadre)
        return pic

    # --- texte ----------------------------------------------------------------------------------
    def texte(self, x, y, w, h, contenu, taille=16, couleur=C["encre"], gras=False, align="l", ancre="t",
              interligne=1.15, espace_apres=0, puces=False):
        """contenu : chaîne, ou liste de paragraphes ; un paragraphe = chaîne ou liste de (texte, gras, couleur, taille)."""
        paragraphes = contenu if isinstance(contenu, list) else [contenu]
        tb = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[ancre]
        lignes_apercu = []
        for i, par in enumerate(paragraphes):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            p.line_spacing = interligne
            p.space_after = Pt(espace_apres)
            morceaux = par if isinstance(par, list) else [(par, gras, couleur, taille)]
            if puces:
                self._puce(p)
            for txt, g, coul, t in morceaux:
                r = p.add_run()
                r.text = txt
                r.font.name, r.font.size, r.font.bold = "Calibri", Pt(t), g
                r.font.color.rgb = RGBColor.from_string(coul)
            lignes_apercu.append((morceaux, puces))
        self._dessiner_texte(x, y, w, h, lignes_apercu, align, ancre, interligne, espace_apres)
        return tb

    @staticmethod
    def _puce(p):
        from pptx.oxml.ns import qn
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.25))))
        pPr.set("indent", str(-Emu(Inches(0.25))))
        for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
            for e in pPr.findall(qn(tag)):
                pPr.remove(e)
        from lxml import etree
        buclr = etree.SubElement(pPr, qn("a:buClr"))
        srgb = etree.SubElement(buclr, qn("a:srgbClr"))
        srgb.set("val", C["orange"])
        bu = etree.SubElement(pPr, qn("a:buChar"))
        bu.set("char", "•")

    def _dessiner_texte(self, x, y, w, h, paragraphes, align, ancre, interligne, espace_apres):
        """Aperçu : coupe les lignes avec les vraies métriques de Calibri et signale les débordements."""
        largeur_px = w * ECHELLE
        blocs = []
        for morceaux, puce in paragraphes:
            retrait = 0.25 * ECHELLE if puce else 0
            mots = []
            for txt, g, coul, t in morceaux:
                police = ImageFont.truetype(FONT[g], int(round(t * ECHELLE / 72)))
                for k, mot in enumerate(txt.split(" ")):
                    mots.append((mot, police, coul, t, k > 0 or (mots and txt.startswith(" "))))
            lignes, courante, largeur = [], [], 0
            for mot, police, coul, t, espace in mots:
                lmot = police.getlength(mot)
                lesp = police.getlength(" ") if courante else 0
                if courante and largeur + lesp + lmot > largeur_px - retrait + 0.5:
                    lignes.append(courante)
                    courante, largeur, lesp = [], 0, 0
                courante.append((mot, police, coul, t, lesp))
                largeur += lesp + lmot
                if lmot > largeur_px - retrait + 0.5:
                    self.deck.alertes.append(f"diapo {len(self.deck.prs.slides)} : mot trop long « {mot} »")
            if courante:
                lignes.append(courante)
            taille_max = max(t for _, _, _, t, _ in mots) if mots else 14
            blocs.append((lignes, taille_max, retrait, puce))
        hauteur = sum(len(l) * tm * interligne * 1.2 * ECHELLE / 72 + espace_apres * ECHELLE / 72
                      for l, tm, _, _ in blocs) - espace_apres * ECHELLE / 72
        if hauteur > h * ECHELLE + 2:
            premier = "".join(m[0] for m in paragraphes[0][0])[:40]
            self.deck.alertes.append(f"diapo {len(self.deck.prs.slides)} : texte trop haut ({hauteur / ECHELLE:.2f} in "
                                     f"pour {h:.2f} in) « {premier}… »")
        y_px = y * ECHELLE + {"t": 0, "m": max(0, (h * ECHELLE - hauteur) / 2), "b": max(0, h * ECHELLE - hauteur)}[ancre]
        for lignes, tm, retrait, puce in blocs:
            hl = tm * interligne * 1.2 * ECHELLE / 72
            for n, ligne in enumerate(lignes):
                lg = sum(m[1].getlength(m[0]) + m[4] for m in ligne)
                x_px = x * ECHELLE + retrait + {"l": 0, "c": (largeur_px - retrait - lg) / 2, "r": largeur_px - retrait - lg}[align]
                if puce and n == 0:
                    self.draw.text((x * ECHELLE, y_px + (hl - tm * ECHELLE / 72) / 2), "•",
                                   font=ligne[0][1], fill="#" + C["orange"])
                for mot, police, coul, t, esp in ligne:
                    x_px += esp
                    self.draw.text((x_px, y_px + (hl - t * ECHELLE / 72) / 2), mot, font=police, fill="#" + coul)
                    x_px += police.getlength(mot)
                y_px += hl
            y_px += espace_apres * ECHELLE / 72

    def titre(self, texte, sous_titre=None):
        self.texte(0.6, 0.45, L - 1.2, 0.75, texte, 34, C["nuit"], gras=True, ancre="m")
        if sous_titre:
            self.texte(0.6, 1.2, L - 1.2, 0.45, sous_titre, 16, C["gris"])

    def carte(self, x, y, w, h, fond=C["clair"]):
        return self.forme(x, y, w, h, fond, "arrondi")

    def pastille(self, x, y, d, contenu, fond=C["orange"], taille=16, couleur=C["blanc"]):
        self.forme(x, y, d, d, fond, "cercle")
        self.texte(x, y, d, d, contenu, taille, couleur, gras=True, align="c", ancre="m")

    def notes(self, texte):
        self.slide.notes_slide.notes_text_frame.text = texte


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(L), Inches(H)
        self.alertes, self.apercus = [], []

    def nouvelle(self, fond=C["blanc"], numero=True):
        n = len(self.prs.slides) + 1
        d = Diapo(self, fond, n if numero else None)
        self.apercus.append(d)
        return d


# =============================================================================================
def construire(d):
    deck = Deck()
    tdb, ir, irec = d["tdb"], d["ind_rh"], d["ind_rec"]
    nb_tests, echecs = d["tests"]
    regles_ok, _ = d["controle"]
    profil = d["profil"]

    # 1. Titre ------------------------------------------------------------------------------------
    s = deck.nouvelle(C["nuit"], numero=False)
    s.forme(9.4, 1.0, 4.6, 4.6, C["nuit_2"], "cercle")
    s.forme(10.9, 3.6, 2.6, 2.6, C["orange"], "cercle")
    s.forme(9.9, 4.9, 1.1, 1.1, C["bleu"], "cercle")
    s.texte(0.8, 1.2, 8.4, 0.4, "NURU HR DIGITAL PORTFOLIO", 16, C["orange"], gras=True)
    s.texte(0.8, 1.75, 8.4, 2.0, "Des outils RH fiables, traçables et responsables", 44, C["blanc"], gras=True)
    s.texte(0.8, 3.85, 8.2, 0.9, "Projet personnel de démonstration · candidature au stage en Ressources Humaines · Goma",
            18, "C9D3E2")
    s.texte(0.8, 5.15, 8.2, 0.45, d["nom"], 20, C["blanc"], gras=True)
    s.texte(0.8, 5.6, 8.2, 0.4, f"{profil['titre']} · {cr.date_fr(d['date_document'])}", 14, "C9D3E2")
    s.texte(0.8, 6.65, 11.8, 0.4, "Données entièrement fictives · projet indépendant de Nuru : ni commandé, ni validé, "
            "ni utilisé par Nuru", 12, "AEB9CC")
    s.notes("Bonjour. Je vous présente un projet personnel que j'ai réalisé pour accompagner ma candidature : des outils "
            "de gestion administrative RH construits sur des données entièrement fictives. Je précise d'emblée qu'il "
            "s'agit d'une démonstration indépendante : Nuru ne l'a ni commandé ni validé. Durée : environ 15 minutes, "
            "puis je peux faire une démonstration en direct.")

    # 2. L'essentiel ---------------------------------------------------------------------------------
    s = deck.nouvelle()
    s.titre("L'essentiel en quatre chiffres", f"Situation fictive au {cr.date_fr(d['jour'])}, calculée automatiquement")
    stats = [(str(tdb["effectif_total"]), "collaborateurs actifs suivis dans un registre unique"),
             (str(len(tdb["contrats_30"])), "CDD à échéance sous 30 jours, signalés à temps"),
             (str(tdb["candidatures_total"]), "candidatures suivies en 8 étapes, décisions humaines"),
             (str(nb_tests), "tests automatisés réussis" if not echecs else "tests automatisés (échecs à corriger)")]
    for i, (val, lib) in enumerate(stats):
        x = 0.6 + i * 3.08
        s.carte(x, 2.05, 2.78, 3.0)
        s.texte(x + 0.3, 2.3, 2.2, 1.3, val, 60, C["orange"], gras=True)
        s.texte(x + 0.3, 3.7, 2.2, 1.2, lib, 15, C["encre"])
    s.forme(0.6, 5.45, 12.13, 1.1, C["nuit"], "arrondi")
    s.texte(1.0, 5.45, 11.4, 1.1, [[("Le principe : ", True, C["orange"], 18),
                                    ("des données fiables, des alertes au bon moment, et des décisions qui restent humaines.",
                                     False, C["blanc"], 18)]], ancre="m")
    s.notes(f"Quatre chiffres résument le projet. {tdb['effectif_total']} collaborateurs actifs, tous fictifs. "
            f"{len(tdb['contrats_30'])} CDD arrivent à échéance dans les 30 jours : l'outil les signale pour que la "
            f"décision soit préparée à temps. {tdb['candidatures_total']} candidatures suivies en 8 étapes. Et "
            f"{nb_tests} tests automatisés qui vérifient chaque calcul. Le fil conducteur : fiabilité, anticipation, "
            "décisions humaines.")

    # 3. Profil ---------------------------------------------------------------------------------------
    s = deck.nouvelle()
    s.titre("Mon profil", "Une formation en management des organisations appliquée à la fonction RH")
    formations = [("M", "Master en Management des organisations"),
                  ("L", "Licence en Sciences administratives et Management"),
                  ("RH", "Formation certifiée : GRH dans le contexte humanitaire"),
                  ("IA", "Formation certifiée : IA appliquée au monde professionnel avec Claude")]
    for i, (sigle, lib) in enumerate(formations):
        y = 2.0 + i * 1.12
        s.pastille(0.6, y, 0.8, sigle, C["nuit"] if i < 2 else C["orange"], 16 if len(sigle) == 1 else 15)
        s.texte(1.65, y, 5.3, 0.8, lib, 17, C["encre"], gras=True, ancre="m")
    s.carte(7.3, 2.0, 5.43, 4.3)
    s.texte(7.7, 2.25, 4.7, 0.5, "Mon approche", 20, C["nuit"], gras=True)
    s.texte(7.7, 2.9, 4.75, 3.3, ["Raisonner en processus, en rôles et en indicateurs",
                                  "Laisser une trace écrite de chaque étape",
                                  "Anticiper plutôt que réagir dans l'urgence",
                                  "Protéger les données des personnes",
                                  "Utiliser l'IA comme assistant, jamais comme décideur"],
            16, C["encre"], puces=True, espace_apres=8)
    s.notes("Ma formation en management des organisations m'a appris à analyser un fonctionnement en processus, en rôles "
            "et en indicateurs. C'est ce regard que j'applique à la gestion administrative RH. Ma formation en GRH dans le "
            "contexte humanitaire m'a sensibilisé(e) à l'équité et à la confidentialité ; ma formation en IA avec Claude "
            "m'a appris à utiliser un assistant de façon encadrée. (Adapter : établissements et années.)")

    # 4. Problèmes ---------------------------------------------------------------------------------------
    s = deck.nouvelle()
    s.titre("Six problèmes RH concrets, six réponses", "Ce que le projet cherche à résoudre")
    for i, (probleme, reponse) in enumerate(cr.PROBLEMES):
        col, lig = i % 3, i // 3
        x, y = 0.6 + col * 4.12, 2.0 + lig * 2.35
        s.carte(x, y, 3.85, 2.05)
        s.pastille(x + 0.3, y + 0.3, 0.6, str(i + 1), C["orange"], 16)
        s.texte(x + 1.1, y + 0.25, 2.5, 0.75, probleme, 16, C["nuit"], gras=True, ancre="m")
        s.texte(x + 0.3, y + 1.1, 3.3, 0.8, reponse, 14, C["encre"])
    s.notes("J'ai identifié six difficultés fréquentes dans la gestion administrative du personnel : des dossiers "
            "difficiles à retrouver, des contrats qui arrivent à échéance sans décision préparée, des effectifs mal "
            "connus, des absences suivies dans des fichiers dispersés, un recrutement peu traçable et l'absence de vue "
            "d'ensemble. Chaque outil répond à l'un de ces problèmes.")

    # 5. Démarche ------------------------------------------------------------------------------------------
    s = deck.nouvelle()
    s.titre("Ma démarche en six étapes", "Une chaîne reproductible : une commande suffit pour tout régénérer")
    etapes = [("Données", f"{d['volumes']['collaborateurs']} collaborateurs fictifs"),
              ("Contrôle", f"{regles_ok} règles de cohérence"), ("Outils", "3 pages de consultation"),
              ("Documents", "procédures, rapport, modèles"), ("Insights", f"{len(d['insights'])} constats chiffrés"),
              ("Tests", f"{nb_tests} tests automatisés")]
    s.trait(1.35, 3.05, 12.0, 3.05, C["trait"], 3)
    for i, (nom, detail) in enumerate(etapes):
        cx = 1.35 + i * 2.13
        s.pastille(cx - 0.5, 2.55, 1.0, str(i + 1), C["nuit"] if i % 2 == 0 else C["orange"], 22)
        s.texte(cx - 1.0, 3.8, 2.0, 0.45, nom, 18, C["nuit"], gras=True, align="c")
        s.texte(cx - 1.0, 4.3, 2.0, 0.8, detail, 14, C["gris"], align="c")
    s.carte(0.6, 5.45, 12.13, 1.05, "EAF1FB")
    s.texte(1.0, 5.45, 11.4, 1.05, "Chaque chiffre affiché est recalculé depuis les fichiers sources : aucune valeur "
            "n'est saisie à la main dans les outils ni dans ce support.", 16, C["nuit"], ancre="m")
    s.notes("La démarche est volontairement simple et reproductible. Je génère des données fictives cohérentes, je les "
            "contrôle avec des règles automatiques, je construis les outils, je rédige les documents, je calcule des "
            "insights, et des tests vérifient l'ensemble. Si une donnée change, tout se recalcule de façon cohérente.")

    # 6. Module de suivi -------------------------------------------------------------------------------------
    s = deck.nouvelle()
    s.titre("Suivi administratif : anticiper les échéances", "Registre, contrats, absences, stagiaires et effectifs sur une seule page")
    s.image(cr.CAPTURES / "module_rh_echeances.png", 0.6, 1.95, w=7.4, h=4.25)
    blocs_stats = [(str(ir["echeances_30"]), "CDD à 30 jours ou moins"),
                   (f"{ir['part_cdd_pct']} %", "de l'effectif en CDD"),
                   (str(ir["stagiaires_en_cours"] + ir["stagiaires_a_venir"]), "stagiaires en cours ou à venir")]
    for i, (val, lib) in enumerate(blocs_stats):
        y = 1.95 + i * 1.1
        s.texte(8.45, y, 1.6, 0.9, val, 36, C["orange"], gras=True, ancre="m")
        s.texte(10.1, y, 2.65, 0.9, lib, 15, C["encre"], ancre="m")
    s.carte(8.45, 5.35, 4.28, 0.85, "EAF1FB")
    s.texte(8.7, 5.35, 3.85, 0.85, "L'outil alerte, le responsable RH décide", 15, C["nuit"], gras=True, ancre="m")
    s.notes("Le module de suivi administratif réunit sur une page le registre du personnel, les contrats, les absences, "
            "les stagiaires et les effectifs. L'onglet Échéances classe les contrats par urgence : 30, 60, 90 jours. "
            "Il n'affiche que le nécessaire : pas de date de naissance, pas de salaire, et les absences pour raison de "
            "santé apparaissent sous une catégorie confidentielle. La décision contractuelle reste celle du responsable RH.")

    # 7. Recrutement -----------------------------------------------------------------------------------------
    s = deck.nouvelle()
    s.titre("Recrutement : 8 étapes, des décisions humaines", "Pour chaque candidature : statut, prochaine action, responsable, échéance")
    noms_etapes = ["Candidature reçue", "Présélection", "Convocation au test", "Test réalisé", "Références",
                   "Décision RH", "Offre et contrat", "Intégration"]
    s.trait(1.1, 2.55, 12.2, 2.55, C["trait"], 3)
    for i, nom in enumerate(noms_etapes):
        cx = 1.1 + i * 1.585
        s.pastille(cx - 0.35, 2.2, 0.7, str(i + 1), C["orange"] if i == 5 else C["nuit"], 16)
        s.texte(cx - 0.78, 3.05, 1.56, 0.7, nom, 13, C["orange"] if i == 5 else C["encre"], gras=i == 5, align="c")
    chiffres_rec = [(str(irec["en_cours"]), "dossiers en cours"), (str(irec["postes_pourvus"]), "poste pourvu"),
                    (f"{irec['delai_moyen_decision_jours']} j", "délai moyen jusqu'à la décision")]
    for i, (val, lib) in enumerate(chiffres_rec):
        x = 0.6 + i * 2.55
        s.carte(x, 4.15, 2.3, 2.1)
        s.texte(x + 0.25, 4.3, 1.9, 0.95, val, 40, C["orange"], gras=True)
        s.texte(x + 0.25, 5.3, 1.9, 0.8, lib, 14, C["encre"])
    s.forme(8.3, 4.15, 4.43, 2.1, C["nuit"], "arrondi")
    s.texte(8.65, 4.3, 3.8, 1.85, ["Aucun score automatique", "Aucun classement des candidats",
                                    "Aucune IA de tri", "Références avec autorisation écrite"],
            15, C["blanc"], puces=True, espace_apres=4, ancre="m")
    s.notes("Le recrutement suit 8 étapes. Le module indique pour chaque candidature la prochaine action, qui doit la "
            "faire et pour quand. Il contrôle le respect du processus : pas de saut d'étape, décision réservée au "
            "responsable RH, références contactées seulement avec l'autorisation du candidat. Il ne note pas et ne "
            "classe pas les candidats, et ne collecte aucune caractéristique sensible.")

    # 8. Tableau de bord --------------------------------------------------------------------------------------
    s = deck.nouvelle()
    s.titre("Tableau de bord RH : piloter en un coup d'œil", "Indicateurs, graphiques, insights et actions à effectuer")
    s.image(cr.CAPTURES / "tableau_de_bord_bureau.png", 0.6, 1.95, w=7.6, h=4.65)
    s.image(cr.CAPTURES / "tableau_de_bord_mobile.png", 8.55, 1.95, w=2.0, h=4.65)
    s.texte(10.85, 1.95, 1.9, 4.65, ["Calculé depuis les données", "Adapté au téléphone", "Export CSV pour Excel",
                                     "Actions triées par urgence"],
            15, C["encre"], puces=True, espace_apres=14, ancre="m")
    s.notes("Le tableau de bord réunit les indicateurs clés, trois graphiques, les insights et la liste des actions à "
            "effectuer, triée par urgence. Il fonctionne sur ordinateur comme sur téléphone, et les indicateurs "
            "s'exportent en CSV pour Excel. Je peux vous le montrer en direct.")

    # 9. Graphique natif : effectifs ---------------------------------------------------------------------------
    s = deck.nouvelle()
    s.titre(f"Effectifs au {cr.date_fr(d['jour'])}", "Répartition par département et type de contrat")
    deps = tdb["effectifs_departements"]
    donnees_graph = CategoryChartData()
    donnees_graph.categories = list(reversed(list(deps)))
    donnees_graph.add_series("CDI", [deps[k]["CDI"] for k in reversed(list(deps))])
    donnees_graph.add_series("CDD", [deps[k]["CDD"] for k in reversed(list(deps))])
    gf = s.slide.shapes.add_chart(XL_CHART_TYPE.BAR_STACKED, Inches(0.6), Inches(1.9), Inches(7.6), Inches(4.8),
                                  donnees_graph)
    ch = gf.chart
    ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.TOP, False
    ch.legend.font.size, ch.legend.font.name = Pt(14), "Calibri"
    ch.font.name, ch.font.size = "Calibri", Pt(13)
    for serie, couleur in zip(ch.series, (C["bleu"], C["orange"])):
        serie.format.fill.solid()
        serie.format.fill.fore_color.rgb = RGBColor.from_string(couleur)
        serie.format.line.color.rgb = RGBColor.from_string(C["blanc"])
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 60, 100
    plot.has_data_labels = True
    plot.data_labels.font.size, plot.data_labels.font.color.rgb = Pt(12), RGBColor.from_string(C["blanc"])
    plot.data_labels.position = XL_LABEL_POSITION.CENTER
    plot.data_labels.show_value = True
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.category_axis.tick_labels.font.size = Pt(13)
    ch.category_axis.format.line.color.rgb = RGBColor.from_string(C["trait"])
    # Aperçu du graphique (barres empilées dessinées à l'échelle)
    maxi = max(v["CDI"] + v["CDD"] for v in deps.values())
    for i, (dep, v) in enumerate(deps.items()):
        yb = 2.45 + i * 0.52
        s.draw.text((0.65 * ECHELLE, yb * ECHELLE), dep, font=ImageFont.truetype(FONT[False], 17), fill="#" + C["encre"])
        x0 = 3.6
        for val, coul in ((v["CDI"], C["bleu"]), (v["CDD"], C["orange"])):
            lw = 4.4 * val / maxi
            if val:
                s.draw.rectangle([x0 * ECHELLE, yb * ECHELLE, (x0 + lw) * ECHELLE, (yb + 0.32) * ECHELLE], fill="#" + coul)
            x0 += lw
    insight_cdd = next((i for i in d["insights"] if i["theme"] == "Effectifs" and "CDD" in i["titre"]), None)
    s.carte(8.6, 1.9, 4.13, 4.8)
    s.texte(8.95, 2.15, 3.5, 0.5, "À retenir", 20, C["nuit"], gras=True)
    s.texte(8.95, 2.75, 3.5, 1.0, f"{tdb['part_cdd_pct']} %", 44, C["orange"], gras=True)
    s.texte(8.95, 3.75, 3.5, 0.5, "de l'effectif en CDD", 15, C["encre"])
    if insight_cdd:
        s.texte(8.95, 4.4, 3.5, 2.1, insight_cdd["constat"].split(" Calcul")[0], 14, C["encre"])
    s.notes(f"Voici la répartition des {tdb['effectif_total']} collaborateurs actifs par département et type de contrat. "
            f"{tdb['part_cdd_pct']} % de l'effectif est en CDD, avec une part plus élevée dans un département : "
            "c'est un point à examiner avec le responsable du département, sans conclusion hâtive. Ce graphique est "
            "natif dans PowerPoint : ses données restent modifiables.")

    # 10. Insights -----------------------------------------------------------------------------------------------
    s = deck.nouvelle()
    s.titre("Ce que disent les données", "Constats agrégés, calculés automatiquement ; la décision reste humaine")
    couleurs_niveau = {cr.ins.ATTENTION: C["orange"], cr.ins.A_SUIVRE: C["bleu"], cr.ins.INFORMATION: C["gris"]}
    for i, it in enumerate(d["insights"][:4]):
        col, lig = i % 2, i // 2
        x, y = 0.6 + col * 6.17, 1.95 + lig * 2.35
        s.carte(x, y, 5.96, 2.1)
        s.forme(x + 0.3, y + 0.3, 1.45, 0.4, couleurs_niveau[it["niveau"]], "arrondi")
        s.texte(x + 0.3, y + 0.3, 1.45, 0.4, it["niveau"], 13, C["blanc"], gras=True, align="c", ancre="m")
        s.texte(x + 0.3, y + 0.85, 1.9, 0.9, it["indicateur"], 30, C["nuit"], gras=True, ancre="m")
        s.texte(x + 2.3, y + 0.25, 3.4, 1.6, it["titre"], 16, C["encre"], gras=True, ancre="m")
    s.texte(0.6, 6.55, 12.1, 0.35, "Pistes proposées (exemples) : préparer les dossiers de fin de CDD, traiter les actions en "
            "retard, comprendre les causes des départs.", 13, C["gris"])
    s.notes("Le tableau de bord calcule aussi des insights : des constats chiffrés, toujours agrégés, sans jugement sur "
            "une personne. " + " ".join(f"{i['titre']} : {i['indicateur']}." for i in d["insights"][:4])
            + " Chaque constat est accompagné d'une piste, mais la décision appartient au responsable RH.")

    # 11. IA responsable -------------------------------------------------------------------------------------------
    s = deck.nouvelle()
    s.titre("IA responsable et confidentialité", "L'IA m'a assisté(e) ; elle ne décide de rien")
    for i, (titre_ia, texte_ia) in enumerate(cr.IA_RESPONSABLE):
        x = 0.6 + i * 4.12
        s.carte(x, 1.95, 3.85, 4.5)
        s.pastille(x + 0.35, 2.25, 0.85, str(i + 1), C["nuit"] if i != 1 else C["orange"], 22)
        s.texte(x + 0.35, 3.3, 3.2, 0.8, titre_ia, 19, C["nuit"], gras=True)
        s.texte(x + 0.35, 4.3, 3.2, 2.0, texte_ia, 14, C["encre"])
    s.notes("J'ai construit ce projet avec l'assistance de Claude Code, et je le dis clairement. Les propositions de l'IA "
            "ont été exécutées et vérifiées par des tests. Trois principes : transparence sur l'usage de l'IA, aucune "
            "décision automatisée sur les personnes, et protection des données dès la conception.")

    # 12. Qualité ----------------------------------------------------------------------------------------------------
    s = deck.nouvelle()
    s.titre("Qualité : ce qui a été réellement testé", "Des vérifications exécutées, pas seulement annoncées")
    qualite = [(str(nb_tests), "tests automatisés réussis" if not echecs else "tests, échecs à corriger"),
               (str(regles_ok), "règles de cohérence respectées"), ("11", "pages vérifiées : liens et téléphone")]
    for i, (val, lib) in enumerate(qualite):
        y = 1.95 + i * 1.5
        s.carte(0.6, y, 5.6, 1.3)
        s.texte(0.95, y, 1.9, 1.3, val, 44, C["orange"], gras=True, ancre="m")
        s.texte(2.9, y, 3.1, 1.3, lib, 16, C["encre"], ancre="m")
    s.forme(6.6, 1.95, 6.13, 4.3, C["nuit"], "arrondi")
    s.texte(6.95, 2.2, 5.5, 0.5, "Limites assumées", 20, C["orange"], gras=True)
    s.texte(6.95, 2.85, 5.45, 3.3, ["Consultation seule, pas de saisie en ligne", "Pas de gestion des droits d'accès",
                                    "Soldes de congés non calculés", "Règles juridiques à faire vérifier"],
            16, C["blanc"], puces=True, espace_apres=10)
    s.notes(f"La qualité a été vérifiée concrètement : {nb_tests} tests automatisés, {regles_ok} règles de cohérence et "
            "11 pages contrôlées, y compris sur une largeur de téléphone. Je présente aussi les limites : ce sont des "
            "outils de consultation, sans gestion des droits ; les soldes de congés et les règles juridiques demandent "
            "une validation humaine.")

    # 13. Et ensuite ----------------------------------------------------------------------------------------------------
    s = deck.nouvelle()
    s.titre("Et ensuite ?", "Ce que j'améliorerais, et ce que j'apporterais à l'équipe RH")
    s.carte(0.6, 1.95, 5.96, 4.5)
    s.texte(0.95, 2.2, 5.3, 0.5, "Pistes d'amélioration", 20, C["nuit"], gras=True)
    s.texte(0.95, 2.85, 5.3, 3.4, cr.PROCHAINES_ETAPES, 15, C["encre"], puces=True, espace_apres=10)
    s.forme(6.77, 1.95, 5.96, 4.5, C["nuit"], "arrondi")
    s.texte(7.12, 2.2, 5.3, 0.5, "Ce que j'apporterais au stage", 20, C["orange"], gras=True)
    s.texte(7.12, 2.85, 5.3, 3.5, [a.split(" : ")[0] if " : " in a else a for a in cr.APPORTS_STAGE],
            15, C["blanc"], puces=True, espace_apres=10)
    s.notes("Pour aller plus loin, j'ajouterais une saisie sécurisée, le calcul des soldes de congés une fois la règle "
            "validée, et une version Power BI. Surtout, j'aimerais apprendre les pratiques réelles de votre équipe : "
            "ce projet montre ma façon de travailler, que je souhaite mettre à votre service et enrichir à votre contact.")

    # 14. Merci ------------------------------------------------------------------------------------------------------------
    s = deck.nouvelle(C["nuit"], numero=False)
    s.forme(9.4, 1.0, 4.6, 4.6, C["nuit_2"], "cercle")
    s.forme(10.9, 3.6, 2.6, 2.6, C["orange"], "cercle")
    s.forme(9.9, 4.9, 1.1, 1.1, C["bleu"], "cercle")
    s.texte(0.8, 1.3, 8.4, 1.2, "Merci", 54, C["blanc"], gras=True)
    s.texte(0.8, 2.5, 8.4, 0.6, "Questions et démonstration en direct", 22, C["orange"])
    contacts = [("Nom : ", d["nom"]), ("E-mail : ", profil["email"] or "à compléter"),
                ("LinkedIn : ", profil["linkedin"] or "à compléter"),
                ("Portfolio en ligne : ", profil["portfolio_existant"] or "lien à compléter après publication")]
    s.texte(0.8, 3.55, 8.4, 2.3, [[(k, True, C["blanc"], 16), (v, False, "C9D3E2", 16)] for k, v in contacts],
            espace_apres=8)
    s.texte(0.8, 6.65, 11.8, 0.4, "Données entièrement fictives · projet personnel indépendant de Nuru", 12, "AEB9CC")
    s.notes("Merci pour votre attention. Je peux maintenant répondre à vos questions ou ouvrir le tableau de bord et les "
            "modules pour une démonstration en direct, sur ordinateur ou sur téléphone.")
    return deck


if __name__ == "__main__":
    print("Collecte des données et exécution des tests…")
    donnees = cr.collecter()
    deck = construire(donnees)
    props = deck.prs.core_properties
    props.title = "NURU HR DIGITAL PORTFOLIO - Présentation"
    props.author = donnees["profil"]["nom_complet"] or "Candidat(e) au stage RH"
    props.subject = "Projet personnel de démonstration RH (données fictives, indépendant de Nuru)"
    deck.prs.save(FICHIER)
    APERCU.mkdir(exist_ok=True)
    for i, diapo in enumerate(deck.apercus, start=1):
        diapo.img.save(APERCU / f"diapo_{i:02d}.png")
    print(f"Présentation générée : {FICHIER.relative_to(cr.RACINE)} ({len(deck.prs.slides)} diapositives)")
    print(f"Aperçus : {APERCU.relative_to(cr.RACINE)}")
    print("Contrôle de mise en page : " + ("aucun débordement détecté" if not deck.alertes else ""))
    for a in deck.alertes:
        print("  ALERTE " + a)
