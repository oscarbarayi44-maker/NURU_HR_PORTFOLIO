"""
Génère le rapport de projet complet en deux formats, à partir du même contenu :
- documents/rapports/Rapport_projet_NURU_HR_Portfolio.docx  (Word)
- documents/rapports/Rapport_projet_NURU_HR_Portfolio.pdf   (PDF)

Utilisation (depuis la racine du projet) :
    python documents/rapports/generer_rapport_complet.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import contenu_rapport as cr  # noqa: E402
import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib import font_manager  # noqa: E402

NOM_FICHIER = "Rapport_projet_NURU_HR_Portfolio"
C = cr.COULEURS


# =============================================================================================
# 1. Graphiques (images partagées par le Word et le PDF)
# =============================================================================================
def _style():
    polices = {f.name for f in font_manager.fontManager.ttflist}
    plt.rcParams.update({
        "font.family": "Calibri" if "Calibri" in polices else "DejaVu Sans", "font.size": 11,
        "axes.spines.top": False, "axes.spines.right": False, "axes.spines.left": False,
        "axes.edgecolor": "#C3C2B7", "axes.labelcolor": "#52514E", "xtick.color": "#52514E",
        "ytick.color": "#1F2933", "axes.grid": True, "axes.grid.axis": "x", "grid.color": "#E1E0D9",
        "grid.linewidth": 0.8, "axes.axisbelow": True,
    })


def _barres_empilees(fichier, libelles, series, titre_legende=None):
    """Barres horizontales empilées, 2 px d'écart blanc entre segments, total en bout de barre."""
    fig, ax = plt.subplots(figsize=(7.2, 0.42 * len(libelles) + 1.0), dpi=200)
    gauche = [0] * len(libelles)
    for nom, valeurs, couleur in series:
        ax.barh(libelles, valeurs, left=gauche, color=f"#{couleur}", height=0.55, label=nom,
                edgecolor="white", linewidth=1.5)
        gauche = [g + v for g, v in zip(gauche, valeurs)]
    for i, total in enumerate(gauche):
        ax.text(total + max(gauche) * 0.015, i, str(total), va="center", fontsize=10.5, color="#1F2933",
                fontweight="bold")
    ax.invert_yaxis()
    ax.set_xlim(0, max(gauche) * 1.12 or 1)
    ax.tick_params(axis="y", length=0)
    ax.xaxis.set_major_locator(matplotlib.ticker.MaxNLocator(integer=True))
    if len(series) > 1:
        ax.legend(loc="lower right", frameon=False, fontsize=10, title=titre_legende, title_fontsize=10)
    fig.tight_layout()
    fig.savefig(fichier, facecolor="white")
    plt.close(fig)


def generer_figures(d):
    _style()
    cr.FIGURES.mkdir(exist_ok=True)
    tdb = d["tdb"]
    deps = tdb["effectifs_departements"]
    _barres_empilees(cr.FIGURES / "effectifs_departements.png", list(deps),
                     [("CDI", [v["CDI"] for v in deps.values()], C["bleu"]),
                      ("CDD", [v["CDD"] for v in deps.values()], C["orange"])])
    offres = list(tdb["offres"])
    _barres_empilees(cr.FIGURES / "recrutement_etapes.png", [f"{k}. {lib}" for k, lib in cr.rc.ETAPES.items()],
                     [(tdb["offres"][o], [d["entonnoir"][o][k] for k in cr.rc.ETAPES], couleur)
                      for o, couleur in zip(offres, (C["bleu"], C["orange"]))])
    cats = tdb["absences"]["jours_par_categorie"]
    _barres_empilees(cr.FIGURES / "absences_categories.png", list(cats), [("Jours", list(cats.values()), C["bleu"])])

    # Évolution mensuelle de l'effectif (ligne, axe ne partant pas de zéro : signalé dans la légende)
    m = d["mensuel"]
    fig, ax = plt.subplots(figsize=(7.2, 2.8), dpi=200)
    mois = [cr.MOIS_FR[int(x["mois"][5:]) - 1][:4] + "." if len(cr.MOIS_FR[int(x["mois"][5:]) - 1]) > 4
            else cr.MOIS_FR[int(x["mois"][5:]) - 1] for x in m]
    valeurs = [x["effectif_fin"] for x in m]
    ax.plot(mois, valeurs, color=f"#{C['bleu']}", linewidth=2, marker="o", markersize=6,
            markeredgecolor="white", markeredgewidth=1.5)
    for i, v in enumerate(valeurs):
        if i in (0, len(valeurs) - 1) or v != valeurs[i - 1]:
            ax.annotate(str(v), (i, v), textcoords="offset points", xytext=(0, 8), ha="center", fontsize=10,
                        color="#1F2933", fontweight="bold")
    ax.set_ylim(min(valeurs) - 3, max(valeurs) + 2)
    ax.yaxis.set_major_locator(matplotlib.ticker.MaxNLocator(integer=True))
    ax.grid(axis="y", color="#E1E0D9")
    ax.grid(axis="x", visible=False)
    ax.set_ylabel("Effectif fin de mois (axe tronqué)", fontsize=9)
    fig.tight_layout()
    fig.savefig(cr.FIGURES / "effectif_mensuel.png", facecolor="white")
    plt.close(fig)


# =============================================================================================
# 2. Contenu structuré (liste de blocs), identique pour le Word et le PDF
# =============================================================================================
def blocs(d):
    tdb, ir, irec = d["tdb"], d["ind_rh"], d["ind_rec"]
    v = d["volumes"]
    nb_tests, echecs = d["tests"]
    regles_ok, regles_ko = d["controle"]
    nb_tables, nb_requetes = d["sql"]
    B = []

    def h1(t): B.append(("h1", t))
    def h2(t): B.append(("h2", t))
    def p(t): B.append(("p", t))
    def puces(items): B.append(("puces", items))
    def table(entetes, lignes, largeurs): B.append(("table", entetes, lignes, largeurs))
    def image(chemin, legende, largeur_cm=16): B.append(("image", chemin, legende, largeur_cm))
    def encadre(titre, texte): B.append(("encadre", titre, texte))
    def chiffres(liste): B.append(("chiffres", liste))
    def saut(): B.append(("saut",))

    h1("1. Résumé exécutif")
    p("NURU HR DIGITAL PORTFOLIO est un projet personnel qui montre, sur un cas entièrement fictif, comment "
      "j'organiserais la gestion administrative des ressources humaines d'une organisation de 30 collaborateurs "
      "répartis entre Goma, Bukavu et Beni. Il réunit des registres fiables, des alertes d'échéance, un "
      "recrutement traçable, un tableau de bord et des documents administratifs.")
    chiffres([(str(tdb["effectif_total"]), "collaborateurs actifs"),
              (str(len(tdb["contrats_30"])), "CDD à échéance sous 30 jours"),
              (str(tdb["candidatures_total"]), "candidatures suivies en 8 étapes"),
              (str(nb_tests), "tests automatisés réussis" if not echecs else "tests (des échecs existent)")])
    puces([
        "Trois outils consultables dans un navigateur, sans installation : module de suivi administratif, "
        "module de recrutement et tableau de bord RH.",
        f"Des chiffres toujours cohérents : {regles_ok} règles de cohérence contrôlées automatiquement, "
        f"{nb_tests} tests exécutés lors de la génération de ce rapport.",
        "Des décisions qui restent humaines : les alertes préparent le travail, le responsable RH décide.",
        "Une protection des données intégrée dès la conception : données fictives, informations minimales, "
        "aucun motif de santé affiché.",
    ])
    encadre("Situation de référence", f"Les chiffres de ce rapport décrivent la situation fictive au "
            f"{cr.date_fr(d['jour'])}, calculée automatiquement à partir des fichiers sources.")

    saut()
    h1("2. Contexte, objectifs et périmètre")
    p("Ce projet accompagne ma candidature au stage en Ressources Humaines à Goma. Il ne décrit aucune "
      "expérience professionnelle : il démontre une méthode de travail et des compétences, sur des données "
      "inventées pour l'occasion.")
    h2("2.1 Les compétences visées")
    table(["Compétence", "Ce que le projet montre"], [list(x) for x in cr.COMPETENCES_VISEES], [6.0, 10.5])
    h2("2.2 Les problèmes RH ciblés")
    table(["Problème fréquent", "Réponse apportée"], [list(x) for x in cr.PROBLEMES], [6.5, 10.0])
    encadre("Indépendance", cr.MENTION)

    h1("3. Démarche et méthodologie")
    p("Le projet suit une chaîne simple et reproductible : chaque étape s'appuie sur la précédente, et une seule "
      "commande suffit pour tout régénérer.")
    table(["Étape", "Contenu", "Contrôle"], [
        ["1. Données fictives", f"{v['collaborateurs']} collaborateurs, {v['contrats']} contrats, {v['absences']} absences, "
                                f"{v['stagiaires']} stagiaires, {v['candidatures']} candidatures", "Générées par script, reproductibles"],
        ["2. Préparation", "Tables enrichies : échéances, registre mensuel, entonnoir", "Recoupées avec SQL"],
        ["3. Contrôle de cohérence", f"{regles_ok + regles_ko} règles (contrats continus, effectifs équilibrés…)",
         f"{regles_ok} respectées" + (f", {regles_ko} en erreur" if regles_ko else "")],
        ["4. Outils", "Suivi administratif, recrutement, tableau de bord", "Vérifiés dans le navigateur"],
        ["5. Documents", "Procédures, rapport mensuel, modèles", "Chiffres issus des données"],
        ["6. Tests", f"{nb_tests} tests automatisés", "Exécutés à chaque génération"],
    ], [3.8, 7.7, 5.0])
    h2("3.1 Outils utilisés")
    table(["Outil", "Utilisation réelle dans le projet"], [list(x) for x in cr.OUTILS], [3.5, 13.0])

    h1("4. Les données fictives")
    p(f"Toutes les données sont synthétiques. Les noms sont des combinaisons inventées ; les adresses e-mail "
      f"utilisent un domaine réservé qui ne peut correspondre à aucune vraie adresse. La base SQL compte "
      f"{nb_tables} tables et {nb_requetes} requêtes de consultation.")
    table(["Fichier", "Lignes", "Contenu"], [
        ["collaborateurs.csv", v["collaborateurs"], "Registre du personnel"],
        ["contrats.csv", v["contrats"], "Historique des contrats (CDD, renouvellements, CDI)"],
        ["absences.csv", v["absences"], "Absences et demandes de congé de 2026"],
        ["stagiaires.csv", v["stagiaires"], "Conventions de stage et encadreurs"],
        ["offres_emploi.csv", v["offres_emploi"], "Offres publiées"],
        ["candidatures.csv", v["candidatures"], "Candidatures, sans caractéristique sensible"],
        ["recrutement_etapes.csv", v["recrutement_etapes"], "Journal des étapes enregistrées par des personnes"],
        ["planning_entretiens.csv", v["planning_entretiens"], "Créneaux de tests et d'entretiens"],
    ], [5.0, 2.0, 9.5])

    saut()
    h1("5. Module de suivi administratif")
    p("Le module rassemble en une page le registre du personnel, les contrats, les échéances, les absences, "
      "les stagiaires et les effectifs. Il affiche uniquement les informations nécessaires au suivi.")
    image(cr.CAPTURES / "module_rh_echeances.png", "Onglet Échéances : alertes graduées (données fictives)")
    puces([
        f"Alertes de contrat : {ir['echeances_30']} CDD à 30 jours ou moins, {ir['echeances_60']} entre 31 et 60 jours, "
        f"{ir['echeances_90']} entre 61 et 90 jours ; {ir['echus']} contrat échu non régularisé.",
        f"Effectif actif : {ir['effectif_actif']} personnes, dont {ir['part_cdd_pct']} % en CDD ; "
        f"{ir['entrees_annee']} entrées et {ir['sorties_annee']} sorties depuis janvier.",
        f"Absences : {cr.nb(ir['jours_absence_realises'])} jours ouvrables réalisés ; {ir['demandes_en_attente']} demande en attente.",
        f"Stagiaires : {ir['stagiaires_en_cours']} en cours et {ir['stagiaires_a_venir']} à venir, chacun avec un encadreur.",
    ])
    image(cr.FIGURES / "effectif_mensuel.png", "Effectif en fin de mois, janvier à octobre 2026 (axe vertical tronqué)")

    h1("6. Module de recrutement")
    p("Chaque candidature suit 8 étapes. Le module indique la prochaine action, la personne responsable et "
      "l'échéance. Il ne note pas, ne classe pas et ne sélectionne pas les candidats : il vérifie seulement que "
      "le processus est respecté.")
    table(["Étape", "Enregistrée par"], [
        ["1. Candidature reçue", "Chargé(e) RH"], ["2. Présélection", "Comité de recrutement (grille)"],
        ["3. Convocation au test", "Chargé(e) RH (préavis d'au moins 5 jours)"],
        ["4. Test réalisé", "Comité de recrutement (grille)"],
        ["5. Vérification des références", "Chargé(e) RH, avec l'autorisation écrite du candidat"],
        ["6. Décision du responsable RH", "Responsable RH uniquement"],
        ["7. Offre et contractualisation", "Responsable RH"], ["8. Intégration", "Chargé(e) RH (checklist)"],
    ], [6.5, 10.0])
    image(cr.FIGURES / "recrutement_etapes.png", "Candidatures ayant franchi chaque étape, par offre")
    p(f"Au {cr.date_fr(d['jour'])} : {irec['en_cours']} dossiers en cours, {irec['actions_a_mener']} actions à mener, "
      f"{irec['postes_pourvus']} poste pourvu ; délai moyen de {irec['delai_moyen_decision_jours']} jours entre la "
      "candidature et la décision du responsable RH. Le module génère aussi le planning des entretiens, les "
      "courriels de convocation et une checklist d'intégration datée.")

    saut()
    h1("7. Tableau de bord RH")
    p("Le tableau de bord réunit les indicateurs clés, trois graphiques, les insights et la liste consolidée des "
      "actions à effectuer, triée par urgence. Il s'adapte au téléphone et permet d'exporter les indicateurs et "
      "les actions en CSV.")
    image(cr.CAPTURES / "tableau_de_bord_bureau.png", "Tableau de bord RH sur ordinateur (données fictives)")
    image(cr.FIGURES / "effectifs_departements.png", "Effectifs actifs par département et type de contrat", 14)
    image(cr.FIGURES / "absences_categories.png", "Jours ouvrables d'absence par catégorie depuis janvier "
                                                  "(motifs de santé regroupés et confidentiels)", 14)

    h1("8. Insights : ce que disent les données")
    p("Ces constats sont calculés automatiquement et présentés de façon agrégée, sans jugement sur une personne. "
      "Les pistes sont des suggestions : la décision appartient au responsable RH.")
    table(["Niveau", "Constat", "Valeur", "Piste d'action"],
          [[i["niveau"], f"{i['titre']}. {i['constat']}", i["indicateur"], i["piste"]] for i in d["insights"]],
          [2.2, 7.3, 2.0, 5.0])

    h1("9. Procédures et documents produits")
    table(["Document", "Objet"], [
        ["PROC-RH-01 · Gestion des congés et des absences", "Demande, validation, registre, confidentialité médicale"],
        ["PROC-RH-02 · Suivi des contrats et des échéances", "Alertes à 90, 60, 30 jours ; décision du responsable RH"],
        ["PROC-RH-03 · Archivage des dossiers individuels", "Plan de classement, nommage, sécurité, conservation"],
        ["Rapport mensuel RH (octobre 2026)", "Généré automatiquement à partir des données"],
        ["Planning des tests et entretiens", "Créneaux, lieux, panel, rappels d'équité"],
        ["Courriel de convocation", "Modèle rempli à partir du planning"],
        ["Checklist d'intégration", "Tâches datées de J-10 à J+30"],
    ], [7.5, 9.0])

    saut()
    h1("10. IA responsable et confidentialité")
    for titre, texte in cr.IA_RESPONSABLE:
        encadre(titre, texte)
    puces([
        "Les références d'un candidat ne sont contactées qu'avec son autorisation écrite.",
        "Aucune caractéristique sensible (genre, âge, origine, santé, situation familiale) n'est collectée sur les candidats.",
        "Les règles juridiques ne sont jamais affirmées par l'outil : elles sont signalées « à vérifier ».",
    ])

    h1("11. Qualité : ce qui a été testé")
    chiffres([(str(nb_tests), "tests automatisés" + (" réussis" if not echecs else f", {echecs} échec(s)")),
              (str(regles_ok), "règles de cohérence respectées"),
              ("11", "pages vérifiées : liens et affichage mobile")])
    puces([
        "Chaque indicateur est recalculé indépendamment, directement depuis les fichiers sources ou par SQL.",
        "Les transitions du recrutement sont testées, y compris les cas interdits (saut d'étape, décision "
        "prise par une autre personne que le responsable RH, références sans autorisation).",
        "Vérifié dans Microsoft Edge : affichage des outils, largeur de téléphone (390 px), exports CSV, infobulles.",
    ])
    h2("11.1 Limites connues")
    puces(cr.LIMITES)

    h1("12. Prochaines étapes et conclusion")
    h2("12.1 Améliorations possibles")
    puces(cr.PROCHAINES_ETAPES)
    h2("12.2 Ce que j'apporterais au stage")
    puces(cr.APPORTS_STAGE)
    p("Ce projet n'a pas vocation à remplacer les outils ni les pratiques d'une équipe RH : il montre ma façon de "
      "travailler, que je souhaite mettre au service de l'équipe et enrichir à son contact.")

    h1("Annexe : relancer le projet")
    table(["Commande", "Effet"], [
        ["python analysis/01_generer_donnees_fictives.py", "Crée les données fictives"],
        ["python analysis/03_verifier_coherence.py", "Contrôle de cohérence"],
        ["python dashboard/generer_tableau_de_bord.py", "Tableau de bord et exports CSV"],
        ["python documents/rapports/generer_rapport_complet.py", "Ce rapport (Word et PDF)"],
        ["python documents/rapports/generer_presentation.py", "La présentation PowerPoint"],
        ["python -m unittest discover -s analysis", "Tests automatisés"],
    ], [9.5, 7.0])
    return B


# =============================================================================================
# 3. Word (python-docx)
# =============================================================================================
def ecrire_word(d, liste, chemin):
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    rgb = lambda h: RGBColor.from_string(h)  # noqa: E731
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name, normal.font.size, normal.font.color.rgb = "Calibri", Pt(11), rgb(C["encre"])
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for nom, taille, couleur, avant in (("Heading 1", 18, C["nuit"], 18), ("Heading 2", 13, C["orange"], 12)):
        s = styles[nom]
        s.font.name, s.font.size, s.font.bold, s.font.color.rgb = "Calibri", Pt(taille), True, rgb(couleur)
        rf = s.element.rPr.find(qn("w:rFonts"))  # supprime les polices du thème pour imposer Calibri
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf is not None and rf.get(qn(att)):
                del rf.attrib[qn(att)]
        s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(avant), Pt(6)
        s.paragraph_format.keep_with_next = True

    def ombrer(cellule, couleur):
        tcPr = cellule._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear"), shd.set(qn("w:color"), "auto"), shd.set(qn("w:fill"), couleur)
        tcPr.append(shd)

    def inserer(tblPr, element, avant):
        """Insère dans tblPr en respectant l'ordre imposé par le schéma Word."""
        for nom in avant:
            suivant = tblPr.find(qn(nom))
            if suivant is not None:
                suivant.addprevious(element)
                return
        tblPr.append(element)

    def bordures(table, couleur=C["trait"]):
        tblPr = table._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{cote}")
            e.set(qn("w:val"), "single" if cote in ("top", "bottom", "insideH") else "nil")
            e.set(qn("w:sz"), "4"), e.set(qn("w:color"), couleur)
            b.append(e)
        inserer(tblPr, b, ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"))

    def marges_cellules(table, haut=80, cote=110):
        tblPr = table._tbl.tblPr
        m = OxmlElement("w:tblCellMar")
        for nom, val in (("top", haut), ("left", cote), ("bottom", haut), ("right", cote)):  # ordre du schéma
            e = OxmlElement(f"w:{nom}")
            e.set(qn("w:w"), str(val)), e.set(qn("w:type"), "dxa")
            m.append(e)
        inserer(tblPr, m, ("w:tblLook", "w:tblCaption", "w:tblDescription"))

    def champ(paragraphe, code):
        run = paragraphe.add_run()
        for typ, txt in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
            if typ:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), typ)
                run._r.append(e)
            else:
                e = OxmlElement("w:instrText" if txt == code else "w:t")
                e.set(qn("xml:space"), "preserve")
                e.text = txt
                run._r.append(e)
        return run

    # --- Page de garde (pas d'en-tête ni de pied de page) ---------------------------------
    sec.different_first_page_header_footer = True
    bandeau = doc.add_table(rows=1, cols=1)
    bandeau.alignment = WD_TABLE_ALIGNMENT.CENTER
    cel = bandeau.cell(0, 0)
    cel.width = Cm(16.6)
    ombrer(cel, C["nuit"])
    marges_cellules(bandeau, 500, 400)
    def ligne(texte, taille, couleur, gras=False, apres=6, italique=False):
        par = cel.add_paragraph() if cel.paragraphs[0].text or len(cel.paragraphs) > 1 else cel.paragraphs[0]
        r = par.add_run(texte)
        r.font.size, r.font.bold, r.font.italic, r.font.color.rgb = Pt(taille), gras, italique, rgb(couleur)
        par.paragraph_format.space_after = Pt(apres)
        return par
    ligne("NURU HR DIGITAL PORTFOLIO", 12, C["orange"], True, 24)
    ligne("Rapport de projet", 30, C["blanc"], True, 4)
    ligne("Des outils de gestion administrative RH fiables, traçables et responsables", 15, "DCE3EE", False, 30)
    ligne(d["nom"], 14, C["blanc"], True, 2)
    ligne(d["profil"]["titre"] + " · " + d["profil"]["ville"], 11, "C9D3E2", False, 2)
    ligne(f"Document du {cr.date_fr(d['date_document'])} · données fictives au {cr.date_fr(d['jour'])}", 11, "C9D3E2", False, 6)
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    avert = doc.add_table(rows=1, cols=1)
    ombrer(avert.cell(0, 0), "FDF0E8")
    marges_cellules(avert, 160, 240)
    pa = avert.cell(0, 0).paragraphs[0]
    r = pa.add_run("Important : ")
    r.bold, r.font.color.rgb = True, rgb(C["orange"])
    pa.add_run(cr.MENTION + " Ce document ne décrit aucune expérience professionnelle.")
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    sommaire = doc.add_paragraph()
    rs = sommaire.add_run("Sommaire")
    rs.bold, rs.font.size, rs.font.color.rgb = True, Pt(14), rgb(C["nuit"])
    for bloc in liste:
        if bloc[0] == "h1":
            ps = doc.add_paragraph(bloc[1])
            ps.paragraph_format.space_after = Pt(2)
            ps.paragraph_format.left_indent = Cm(0.3)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # --- En-tête et pied de page des pages suivantes ----------------------------------------
    ent = sec.header.paragraphs[0]
    ent.text = "NURU HR DIGITAL PORTFOLIO · Rapport de projet · données fictives"
    ent.runs[0].font.size, ent.runs[0].font.color.rgb = Pt(9), rgb(C["gris"])
    pied = sec.footer.paragraphs[0]
    pied.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp = pied.add_run("Projet personnel indépendant de Nuru · page ")
    rp.font.size, rp.font.color.rgb = Pt(9), rgb(C["gris"])
    champ(pied, "PAGE").font.size = Pt(9)

    # --- Corps --------------------------------------------------------------------------------
    for bloc in liste:
        genre = bloc[0]
        if genre == "h1":
            doc.add_heading(bloc[1], level=1)
        elif genre == "h2":
            doc.add_heading(bloc[1], level=2)
        elif genre == "p":
            doc.add_paragraph(bloc[1])
        elif genre == "puces":
            for item in bloc[1]:
                pp = doc.add_paragraph(item, style="List Bullet")
                pp.paragraph_format.space_after = Pt(3)
        elif genre == "saut":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif genre == "table":
            _, entetes, lignes, largeurs = bloc
            t = doc.add_table(rows=1, cols=len(entetes))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            bordures(t)
            marges_cellules(t)
            for i, e in enumerate(entetes):
                c = t.rows[0].cells[i]
                ombrer(c, C["nuit"])
                r = c.paragraphs[0].add_run(e)
                r.bold, r.font.size, r.font.color.rgb = True, Pt(10), rgb(C["blanc"])
            for n, l in enumerate(lignes):
                cells = t.add_row().cells
                for i, val in enumerate(l):
                    r = cells[i].paragraphs[0].add_run(str(val))
                    r.font.size = Pt(10)
                    if i == 0:
                        r.bold = True
                    if n % 2:
                        ombrer(cells[i], C["clair"])
            for n, row in enumerate(t.rows):
                for i, w in enumerate(largeurs):
                    row.cells[i].width = Cm(w)
                    fmt = row.cells[i].paragraphs[0].paragraph_format
                    fmt.space_after = Pt(0)
                    # Tableau court : chaque ligne reste avec la suivante (pas de coupure de page)
                    fmt.keep_with_next = len(lignes) <= 9 and n < len(t.rows) - 1
                trPr_ligne = row._tr.get_or_add_trPr()
                interdit = OxmlElement("w:cantSplit")
                trPr_ligne.append(interdit)
            # En-tête répété si le tableau passe sur une autre page
            trPr = t.rows[0]._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            th.set(qn("w:val"), "true")
            trPr.append(th)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif genre == "image":
            _, chemin_img, legende, largeur = bloc
            pi = doc.add_paragraph()
            pi.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pi.paragraph_format.keep_with_next = True
            pi.add_run().add_picture(str(chemin_img), width=Cm(largeur))
            pl = doc.add_paragraph(legende)
            pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pl.runs[0].font.size, pl.runs[0].font.italic, pl.runs[0].font.color.rgb = Pt(9), True, rgb(C["gris"])
        elif genre == "encadre":
            _, titre, texte = bloc
            t = doc.add_table(rows=1, cols=1)
            ombrer(t.cell(0, 0), "EAF1FB")
            marges_cellules(t, 120, 200)
            pe = t.cell(0, 0).paragraphs[0]
            r = pe.add_run(titre + " : ")
            r.bold, r.font.color.rgb = True, rgb(C["nuit"])
            pe.add_run(texte)
            pe.paragraph_format.space_after = Pt(0)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif genre == "chiffres":
            t = doc.add_table(rows=2, cols=len(bloc[1]))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            marges_cellules(t, 120, 120)
            for i, (val, lib) in enumerate(bloc[1]):
                for ligne_t in (0, 1):
                    ombrer(t.cell(ligne_t, i), C["clair"])
                    t.cell(ligne_t, i).width = Cm(16.6 / len(bloc[1]))
                pv = t.cell(0, i).paragraphs[0]
                pv.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rv = pv.add_run(val)
                rv.bold, rv.font.size, rv.font.color.rgb = True, Pt(26), rgb(C["orange"])
                pv.paragraph_format.space_after = Pt(0)
                pl = t.cell(1, i).paragraphs[0]
                pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
                rl = pl.add_run(lib)
                rl.font.size, rl.font.color.rgb = Pt(10), rgb(C["gris"])
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

    zoom = doc.settings.element.find(qn("w:zoom"))   # le modèle par défaut omet un attribut requis
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    props = doc.core_properties
    props.title, props.subject = "NURU HR DIGITAL PORTFOLIO - Rapport de projet", "Projet personnel de démonstration RH"
    props.author = d["profil"]["nom_complet"] or "Candidat(e) au stage RH"
    props.comments = cr.MENTION
    doc.save(chemin)


# =============================================================================================
# 4. PDF (reportlab), même contenu
# =============================================================================================
def ecrire_pdf(d, liste, chemin):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (Image, KeepTogether, ListFlowable, ListItem, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)
    from xml.sax.saxutils import escape

    polices = Path("C:/Windows/Fonts")
    if (polices / "calibri.ttf").exists():
        pdfmetrics.registerFont(TTFont("Calibri", str(polices / "calibri.ttf")))
        pdfmetrics.registerFont(TTFont("Calibri-Bold", str(polices / "calibrib.ttf")))
        pdfmetrics.registerFont(TTFont("Calibri-Italic", str(polices / "calibrii.ttf")))
        pdfmetrics.registerFontFamily("Calibri", normal="Calibri", bold="Calibri-Bold", italic="Calibri-Italic")
        police, gras, italique = "Calibri", "Calibri-Bold", "Calibri-Italic"
    else:
        police, gras, italique = "Helvetica", "Helvetica-Bold", "Helvetica-Oblique"

    col = lambda h: colors.HexColor("#" + h)  # noqa: E731
    st = {
        "p": ParagraphStyle("p", fontName=police, fontSize=10.5, leading=14.5, textColor=col(C["encre"]), spaceAfter=6),
        "h1": ParagraphStyle("h1", fontName=gras, fontSize=17, leading=21, textColor=col(C["nuit"]), spaceBefore=14, spaceAfter=8),
        "h2": ParagraphStyle("h2", fontName=gras, fontSize=12.5, leading=16, textColor=col(C["orange"]), spaceBefore=10, spaceAfter=5),
        "cel": ParagraphStyle("cel", fontName=police, fontSize=9.2, leading=12, textColor=col(C["encre"])),
        "cel_g": ParagraphStyle("celg", fontName=gras, fontSize=9.2, leading=12, textColor=col(C["encre"])),
        "ent": ParagraphStyle("ent", fontName=gras, fontSize=9.2, leading=12, textColor=colors.white),
        "leg": ParagraphStyle("leg", fontName=italique, fontSize=8.5, leading=11, textColor=col(C["gris"]), alignment=TA_CENTER, spaceAfter=10),
        "val": ParagraphStyle("val", fontName=gras, fontSize=24, leading=28, textColor=col(C["orange"]), alignment=TA_CENTER),
        "lib": ParagraphStyle("lib", fontName=police, fontSize=9, leading=11.5, textColor=col(C["gris"]), alignment=TA_CENTER),
    }
    largeur_utile = 21 * cm - 2 * 2.2 * cm
    histoire = []

    def P(texte, style="p"):
        return Paragraph(escape(str(texte)), st[style])

    # --- Page de garde ---------------------------------------------------------------------------
    histoire.append(Spacer(1, 7.2 * cm))
    sommaire = [b[1] for b in liste if b[0] == "h1"]
    histoire.append(Paragraph(f"<b>Important :</b> {escape(cr.MENTION)} Ce document ne décrit aucune expérience "
                              "professionnelle.", ParagraphStyle("av", parent=st["p"], backColor=col("FDF0E8"),
                                                                 borderPadding=8, spaceAfter=18)))
    histoire.append(Spacer(1, 0.3 * cm))
    histoire.append(Paragraph("Sommaire", ParagraphStyle("som", parent=st["h1"], fontSize=14, spaceBefore=6)))
    for s in sommaire:
        histoire.append(Paragraph(escape(s), ParagraphStyle("si", parent=st["p"], leftIndent=8, spaceAfter=2)))
    histoire.append(PageBreak())

    for bloc in liste:
        genre = bloc[0]
        if genre in ("h1", "h2"):
            histoire.append(P(bloc[1], genre))
        elif genre == "p":
            histoire.append(P(bloc[1]))
        elif genre == "puces":
            histoire.append(ListFlowable([ListItem(P(i), leftIndent=12, value="•") for i in bloc[1]],
                                         bulletType="bullet", start="•", leftIndent=12, bulletColor=col(C["orange"]),
                                         spaceAfter=6))
        elif genre == "saut":
            histoire.append(PageBreak())
        elif genre == "table":
            _, entetes, lignes, largeurs = bloc
            total = sum(largeurs)
            data = [[P(e, "ent") for e in entetes]] + [
                [P(v, "cel_g" if i == 0 else "cel") for i, v in enumerate(l)] for l in lignes]
            t = Table(data, colWidths=[largeur_utile * w / total for w in largeurs], repeatRows=1)
            style = [("BACKGROUND", (0, 0), (-1, 0), col(C["nuit"])), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                     ("LINEBELOW", (0, 0), (-1, -1), 0.4, col(C["trait"])),
                     ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                     ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6)]
            style += [("BACKGROUND", (0, r), (-1, r), col(C["clair"])) for r in range(2, len(data), 2)]
            t.setStyle(TableStyle(style))
            # Un tableau court n'est jamais coupé entre deux pages
            histoire += [KeepTogether([t]) if len(lignes) <= 9 else t, Spacer(1, 0.35 * cm)]
        elif genre == "image":
            _, chemin_img, legende, largeur = bloc
            from PIL import Image as PILImage
            with PILImage.open(chemin_img) as im:
                ratio = im.height / im.width
            w = min(largeur * cm, largeur_utile)
            histoire.append(KeepTogether([Image(str(chemin_img), width=w, height=w * ratio), Spacer(1, 3),
                                          P(legende, "leg")]))
        elif genre == "encadre":
            _, titre, texte = bloc
            histoire.append(Paragraph(f"<b>{escape(titre)} :</b> {escape(texte)}",
                                      ParagraphStyle("enc", parent=st["p"], backColor=col("EAF1FB"), borderPadding=7,
                                                     spaceBefore=4, spaceAfter=12)))
        elif genre == "chiffres":
            n = len(bloc[1])
            t = Table([[P(v, "val") for v, _ in bloc[1]], [P(l, "lib") for _, l in bloc[1]]],
                      colWidths=[largeur_utile / n] * n)
            t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), col(C["clair"])),
                                   ("LINEAFTER", (0, 0), (-2, -1), 2, colors.white),
                                   ("TOPPADDING", (0, 0), (-1, 0), 8), ("BOTTOMPADDING", (0, 1), (-1, 1), 8),
                                   ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
            histoire += [t, Spacer(1, 0.35 * cm)]

    def premiere_page(canvas, document):
        canvas.saveState()
        canvas.setFillColor(col(C["nuit"]))
        canvas.rect(0, A4[1] - 9.5 * cm, A4[0], 9.5 * cm, fill=1, stroke=0)
        x, y = 2.2 * cm, A4[1] - 2.6 * cm
        canvas.setFillColor(col(C["orange"]))
        canvas.setFont(gras, 11)
        canvas.drawString(x, y, "NURU HR DIGITAL PORTFOLIO")
        canvas.setFillColor(colors.white)
        canvas.setFont(gras, 30)
        canvas.drawString(x, y - 1.4 * cm, "Rapport de projet")
        canvas.setFillColor(col("DCE3EE"))
        canvas.setFont(police, 13.5)
        canvas.drawString(x, y - 2.3 * cm, "Des outils de gestion administrative RH fiables, traçables et responsables")
        canvas.setFillColor(colors.white)
        canvas.setFont(gras, 13)
        canvas.drawString(x, y - 4.0 * cm, d["nom"])
        canvas.setFillColor(col("C9D3E2"))
        canvas.setFont(police, 10.5)
        canvas.drawString(x, y - 4.6 * cm, f"{d['profil']['titre']} · {d['profil']['ville']}")
        canvas.drawString(x, y - 5.2 * cm, f"Document du {cr.date_fr(d['date_document'])} · données fictives au "
                                           f"{cr.date_fr(d['jour'])}")
        canvas.restoreState()

    def pages_suivantes(canvas, document):
        canvas.saveState()
        canvas.setFont(police, 8.5)
        canvas.setFillColor(col(C["gris"]))
        canvas.drawString(2.2 * cm, A4[1] - 1.3 * cm, "NURU HR DIGITAL PORTFOLIO · Rapport de projet · données fictives")
        canvas.drawRightString(A4[0] - 2.2 * cm, 1.1 * cm, f"Projet personnel indépendant de Nuru · page {document.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(chemin), pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm,
                            topMargin=2.0 * cm, bottomMargin=1.9 * cm,
                            title="NURU HR DIGITAL PORTFOLIO - Rapport de projet",
                            author=d["profil"]["nom_complet"] or "Candidat(e) au stage RH",
                            subject="Projet personnel de démonstration RH (données fictives)")
    doc.build(histoire, onFirstPage=premiere_page, onLaterPages=pages_suivantes)


if __name__ == "__main__":
    print("Collecte des données et exécution des tests…")
    donnees = cr.collecter()
    print(f"  tests : {donnees['tests'][0]} exécutés, {donnees['tests'][1]} échec(s)")
    generer_figures(donnees)
    contenu = blocs(donnees)
    ecrire_word(donnees, contenu, cr.RAPPORTS / f"{NOM_FICHIER}.docx")
    ecrire_pdf(donnees, contenu, cr.RAPPORTS / f"{NOM_FICHIER}.pdf")
    print(f"Rapport généré : documents/rapports/{NOM_FICHIER}.docx et .pdf")
